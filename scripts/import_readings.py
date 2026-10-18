"""Import original DOCX paragraphs and run formatting without rewriting text."""
import json
import re
import argparse
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--source-dir', type=Path, default=Path.home() / 'Downloads')
args = parser.parse_args()
FILES = [('Tình cảm', 'DIM - Project.docx'), ('Gia đình', 'DIM - Project (1).docx'), ('Sự nghiệp', 'DIM - Project (2).docx'), ('Học tập', 'DIM - Project (3).docx')]
IMAGES = [1, 2, 4, 3, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 22, 21]
# Display names follow the labels on the card artwork.
CARD_TITLES = {1: 'Kép Văn', 2: 'Kép Võ', 3: 'Kép Con', 4: 'Kép Rằn', 5: 'Đào Chiến', 6: 'Đào Bi', 7: 'Đào Lẳng', 8: 'Đào Điên', 9: 'Tướng Gộc', 10: 'Tướng Lác', 11: 'Tướng Phiên', 12: 'Lão Văn', 13: 'Lão Võ', 14: 'Lão Tiều', 15: 'Nịnh Gộc', 16: 'Nịnh Mụt', 17: 'Nịnh Chẩu', 18: 'Mụ Quý Tộc', 19: 'Mụ Ác', 20: 'Mụ Tiều', 21: 'Kép Xéo Đỏ', 22: 'Kép Xéo Xanh'}

def inherited(obj, prop, kind):
    value = getattr(getattr(obj, kind), prop)
    if value is not None:
        return value
    style = obj.style
    while style:
        value = getattr(getattr(style, kind), prop)
        if value is not None:
            return value
        style = style.base_style
    return None

def paragraph(p):
    result = {'runs': []}
    for r in p.runs:
        if not r.text:
            continue
        run = {'text': r.text}
        for key in ['bold', 'italic', 'underline']:
            value = inherited(r, key, 'font')
            if value is None:
                style = p.style
                while style and value is None:
                    value = getattr(style.font, key)
                    style = style.base_style
            if value:
                run[key] = True
        result['runs'].append(run)
    assert ''.join(r['text'] for r in result['runs']) == p.text
    for key in ['space_before', 'space_after']:
        value = inherited(p, key, 'paragraph_format')
        if value is not None:
            result[key] = value.pt
    alignment = inherited(p, 'alignment', 'paragraph_format')
    if alignment is not None:
        result['textAlign'] = {0: 'left', 1: 'center', 2: 'right', 3: 'justify'}.get(int(alignment), 'left')
    for key, css in [('first_line_indent', 'textIndent'), ('left_indent', 'paddingLeft'), ('right_indent', 'paddingRight')]:
        value = inherited(p, key, 'paragraph_format')
        if value is not None:
            result[css] = f'{value.pt}pt'
    line = inherited(p, 'line_spacing', 'paragraph_format')
    if line is not None:
        result['lineHeight'] = f'{line.pt}pt' if hasattr(line, 'pt') else line
    if re.match(r'^(I|II|III)\.', p.text):
        result['heading'] = True
    return result

result = {}
for topic, filename in FILES:
    doc = Document(args.source_dir / filename)
    starts = [i for i, p in enumerate(doc.paragraphs) if re.match(r'^LÁ(?: BÀI| SỐ)? \d+', p.text, re.I)]
    assert len(starts) == 22
    cards = []
    for index, start in enumerate(starts):
        end = starts[index + 1] if index + 1 < len(starts) else len(doc.paragraphs)
        paras = doc.paragraphs[start + 1:end]
        marker = next(i for i, p in enumerate(paras) if p.text == 'THÔNG ĐIỆP DÀNH CHO BẠN')
        detail = next(i for i, p in enumerate(paras) if re.match(r'^I\.', p.text))
        title = CARD_TITLES[IMAGES[index]]
        image = next((ROOT / 'public/cards').glob(str(IMAGES[index]) + '. *.webp'))
        sections = [i for i, p in enumerate(paras) if re.match(r'^(I|II|III)\.', p.text)]
        details = {f'detail{i+1}': [paragraph(p) for p in paras[a:(sections[i+1] if i+1 < len(sections) else len(paras))]] for i, a in enumerate(sections)}
        details.setdefault('detail3', [])
        # Check the entire imported body, including empty paragraphs and headings.
        imported = [paragraph(paras[marker]), *[paragraph(p) for p in paras[marker+1:detail]],
                    *[p for section in details.values() for p in section]]
        assert [''.join(r['text'] for r in p['runs']) for p in imported] == [p.text for p in paras[marker:]]
        cards.append({'id': f'{topic.lower().replace(" ", "-")}-{index+1:02}', 'image': '/cards/' + image.name, 'title': title, 'message': '\n'.join(p.text for p in paras[marker+1:detail]), 'messageParagraphs': [paragraph(p) for p in paras[marker+1:detail]], 'messageHeading': paragraph(paras[marker]), **details, 'interpretation': '\n'.join(p.text for p in paras[detail:])})
    result[topic] = cards
shuffle = """export function shuffleDeck(topic) {
  const deck = [...(readingsByTopic[topic] || readingsByTopic["Tình cảm"])];
  for (let index = deck.length - 1; index > 0; index -= 1) {
    const swapIndex = Math.floor(Math.random() * (index + 1));
    [deck[index], deck[swapIndex]] = [deck[swapIndex], deck[index]];
  }
  return deck;
}
"""
(ROOT / 'data/readings.js').write_text('// Original DOCX text and formatting. Regenerate with scripts/import_readings.py.\nexport const readingsByTopic = ' + json.dumps(result, ensure_ascii=False, indent=2) + ';\n\n' + shuffle)
print('Imported', sum(map(len, result.values())), 'cards; paragraph text verified against all four DOCX files.')
